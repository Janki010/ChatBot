from pptx import Presentation

from app.extractors.base import BaseExtractor

class PPTXExtractor(BaseExtractor):
    def extract(self, file_path: str) -> list[dict]:
        results = []

        presentation = Presentation(file_path)

        for slide_number, slide in enumerate(presentation.slides, start=1):

            # Extract text from slide
            text_parts = []

            for shape in slide.shapes:
                if not hasattr(shape, "text"):
                    continue

                text = shape.text.strip()

                if text:
                    text_parts.append(text)

            if text_parts:
                results.append({
                    "source_type": "slide",
                    "source_number": slide_number,
                    "element_type": "text",
                    "text": "\n\n".join(text_parts),
                })

            # Extract tables from slide
            for table_number, shape in enumerate(
                self._get_tables(slide),
                start=1,
            ):
                table_text = self._extract_table(shape)

                if not table_text:
                    continue

                results.append({
                    "source_type": "slide",
                    "source_number": slide_number,
                    "element_type": "table",
                    "metadata": {
                        "table_number": table_number,
                    },
                    "text": table_text,
                })

        return results

    @staticmethod
    def _get_tables(slide):
        return [
            shape
            for shape in slide.shapes
            if shape.has_table
        ]

    @staticmethod
    def _extract_table(table_shape) -> str:
        rows = []

        for row in table_shape.table.rows:
            cells = []

            for cell in row.cells:
                cells.append(cell.text.strip())

            rows.append(" | ".join(cells))

        return "\n".join(rows).strip()