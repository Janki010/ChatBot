from pathlib import Path

from openpyxl import load_workbook
import xlrd

from app.extractors.base import BaseExtractor


class XLSXExtractor(BaseExtractor):
    def extract(self, file_path: str) -> list[dict]:
        extension = Path(file_path).suffix.lower()

        if extension == ".xls":
            return self._extract_xls(file_path)

        if extension == ".xlsx":
            return self._extract_xlsx(file_path)

        raise ValueError(
            f"Unsupported Excel file format: {extension}"
        )

    def _extract_xlsx(self, file_path: str) -> list[dict]:
        results = []

        workbook = load_workbook(
            filename=file_path,
            data_only=True,
            read_only=True,
        )

        try:
            for sheet_number, sheet in enumerate(
                workbook.worksheets,
                start=1,
            ):
                sheet_text = self._extract_openpyxl_sheet(sheet)

                if not sheet_text:
                    continue

                results.append({
                    "source_type": "sheet",
                    "source_number": sheet_number,
                    "element_type": "table",
                    "metadata": {
                        "sheet_name": sheet.title,
                        "max_row": sheet.max_row,
                        "max_column": sheet.max_column,
                    },
                    "text": sheet_text,
                })

        finally:
            workbook.close()

        return results

    @staticmethod
    def _extract_openpyxl_sheet(sheet) -> str:
        rows = []

        for row in sheet.iter_rows(values_only=True):
            values = []

            for value in row:
                if value is None:
                    values.append("")
                else:
                    values.append(str(value).strip())

            if not any(values):
                continue

            rows.append(" | ".join(values))

        return "\n".join(rows).strip()

    def _extract_xls(self, file_path: str) -> list[dict]:
        results = []

        workbook = xlrd.open_workbook(
            file_path,
            on_demand=True,
        )

        try:
            for sheet_number in range(workbook.nsheets):
                sheet = workbook.sheet_by_index(sheet_number)

                sheet_text = self._extract_xlrd_sheet(sheet)

                if not sheet_text:
                    continue

                results.append({
                    "source_type": "sheet",
                    "source_number": sheet_number + 1,
                    "element_type": "table",
                    "metadata": {
                        "sheet_name": sheet.name,
                        "max_row": sheet.nrows,
                        "max_column": sheet.ncols,
                    },
                    "text": sheet_text,
                })

        finally:
            workbook.release_resources()

        return results

    @staticmethod
    def _extract_xlrd_sheet(sheet) -> str:
        rows = []

        for row_number in range(sheet.nrows):
            values = []

            for col_number in range(sheet.ncols):
                value = sheet.cell_value(row_number, col_number)

                if value is None:
                    values.append("")
                else:
                    values.append(str(value).strip())

            if not any(values):
                continue

            rows.append(" | ".join(values))

        return "\n".join(rows).strip()