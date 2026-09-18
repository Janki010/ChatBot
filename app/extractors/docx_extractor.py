from docx import Document

from app.extractors.base import BaseExtractor

class DOCXExtractor(BaseExtractor):
    def extract(self, file_path: str) -> list[dict]:
        results = []

        document = Document(file_path)

        # Extract paragraphs
        for paragraph_number, paragraph in enumerate(
            document.paragraphs,
            start=1,
        ):
            text = paragraph.text.strip()

            if not text:
                continue

            results.append({
                "source_type": "paragraph",
                "source_number": paragraph_number,
                "element_type": "text",
                "metadata": {
                    "style": paragraph.style.name,
                },
                "text": text,
            })

        # Extract tables
        for table_number, table in enumerate(
            document.tables,
            start=1,
        ):
            table_text = self._extract_table(table)

            if not table_text:
                continue

            results.append({
                "source_type": "table",
                "source_number": table_number,
                "element_type": "table",
                "metadata": {
                    "table_number": table_number,
                },
                "text": table_text,
            })

        return results

    @staticmethod
    def _extract_table(table) -> str:
        rows = []

        for row in table.rows:
            cells = []

            for cell in row.cells:
                cells.append(cell.text.strip())

            rows.append(" | ".join(cells))

        return "\n".join(rows).strip()